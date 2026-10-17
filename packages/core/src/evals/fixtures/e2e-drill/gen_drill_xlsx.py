#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gen_drill_xlsx —— BI-21 工位 H 演练夹具:三表账簿 xlsx 生成器(合计/借贷/同比)。

数字与 packages/core/src/evals/fixtures/e2e-drill/drill-data.ts 单一数字源逐字同源:
    合计表  税收收入 65000 + 非税收入 25000 = 合计 90000
    借贷表  ∑借方 = ∑贷方 = 90000
    同比表  (953680-910000)/910000 = 4.8;(65500-62500)/62500 = 4.8

用法:
    python3 gen_drill_xlsx.py <out.xlsx> [--corrupt]
    --corrupt:向合计表注入一处故意错账(合计 90000 → 90050),其余两表保持平账。

退出码:0 = 生成成功;2 = 环境问题(缺 openpyxl)。
"""

from __future__ import annotations

import argparse
import sys

SUBTOTAL_TAMPERED = 90050  # 与 drill-data.ts CORRUPTION.subtotalTampered 同源


def build_workbook(corrupt: bool):
    try:
        from openpyxl import Workbook
    except ImportError:
        sys.stderr.write("[gen-drill-xlsx] 缺少依赖 openpyxl: pip install openpyxl\n")
        raise SystemExit(2) from None

    wb = Workbook()

    # —— 表一:合计(合计校验规则:合计行 = 上方数据行之和) ——
    ws_total = wb.active
    assert ws_total is not None
    ws_total.title = "合计"
    ws_total.append(["项目", "金额(亿元)"])
    ws_total.append(["税收收入", 65000])
    ws_total.append(["非税收入", 25000])
    ws_total.append(["合计", SUBTOTAL_TAMPERED if corrupt else 90000])

    # —— 表二:借贷(借贷平衡规则:∑借方 = ∑贷方) ——
    ws_dc = wb.create_sheet("借贷")
    ws_dc.append(["凭证编号", "摘要", "借方(亿元)", "贷方(亿元)"])
    for row in [
        ("记0001", "拨付转移支付", 12000, 12000),
        ("记0002", "国债发行收款", 8000, 8000),
        ("记0003", "税款入库", 70000, 70000),
    ]:
        ws_dc.append(list(row))

    # —— 表三:同比(同比勾稽规则:同比列 = (本期-上期)/|上期|×100) ——
    ws_yoy = wb.create_sheet("同比")
    ws_yoy.append(["指标", "本期", "上期", "同比(%)"])
    for row in [
        ("国内生产总值(亿元)", 953680, 910000, 4.8),
        ("全社会用电量(亿千瓦时)", 65500, 62500, 4.8),
    ]:
        ws_yoy.append(list(row))

    return wb


def main() -> int:
    parser = argparse.ArgumentParser(prog="gen_drill_xlsx")
    parser.add_argument("output", help="输出 .xlsx 路径")
    parser.add_argument("--corrupt", action="store_true", help="注入合计表故意错账")
    args = parser.parse_args()

    wb = build_workbook(corrupt=args.corrupt)
    wb.save(args.output)
    print(f"[gen-drill-xlsx] 已生成 {args.output}(corrupt={args.corrupt})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
