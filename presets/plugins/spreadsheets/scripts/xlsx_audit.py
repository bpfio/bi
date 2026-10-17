#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""xlsx_audit —— Excel/CSV 账表勾稽审计骨架(BI-18 / spreadsheets 插件)。

规则引擎:三类勾稽规则函数(合计校验 / 借贷平衡 / 同比环比勾稽),
输入统一为二维表(list[list]),输出统一为 findings(list[dict])。
读账依赖:CSV 走标准库 csv;.xlsx 需要 openpyxl(缺库时给出 pip install 提示)。

用法:
    python3 xlsx_audit.py <file.xlsx | file.csv> [--sheet NAME] [--tolerance 0.01] [--json]

退出码:0 = 无 error;1 = 存在 error;2 = 用法/环境问题(缺库、文件不存在)。
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from typing import Any, Callable, Optional

# ---------------------------------------------------------------------------
# 读账:文件 → 二维表
# ---------------------------------------------------------------------------

SUBTOTAL_KEYWORDS = ("合计", "总计", "小计")
TOLERANCE_DEFAULT = 0.01


def _is_subtotal_row(row: list[Any]) -> bool:
    """首格含 合计/总计/小计 即视为小计行。"""
    first = row[0] if row else None
    if not isinstance(first, str):
        return False
    return any(k in first for k in SUBTOTAL_KEYWORDS)


def to_number(value: Any) -> Optional[float]:
    """尽力把单元格值转成 float;失败返回 None(占位符/单位混用由规则自行上报)。"""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("，", "").replace(" ", "")
        if text in ("", "-", "—", "–", "N/A", "n/a"):
            return None
        try:
            return float(text)
        except ValueError:
            return None
    return None


def load_table(path: str, sheet: Optional[str] = None) -> list[list[Any]]:
    """读 .xlsx(需 openpyxl)或 .csv(标准库)为二维表,尽量保留原始类型。"""
    lower = path.lower()
    if lower.endswith(".csv"):
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            return [list(row) for row in csv.reader(handle)]
    if lower.endswith((".xlsx", ".xlsm")):
        return _load_xlsx(path, sheet)
    raise SystemExit(f"[xlsx-audit] 不支持的文件类型: {path}(仅 .xlsx/.xlsm/.csv)")


def _load_xlsx(path: str, sheet: Optional[str] = None) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook  # 延迟导入,避免无谓强依赖
    except ImportError:
        sys.stderr.write(
            "[xlsx-audit] 缺少依赖 openpyxl,读 .xlsx 需要先安装:\n"
            "    pip install openpyxl\n"
            "(或改用 CSV 输入,标准库即可处理)\n"
        )
        raise SystemExit(2) from None
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 —— 把 BadZipFile 等归一为用法错误
        raise SystemExit(f"[xlsx-audit] 无法读取 {path}: {exc}") from exc
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        if worksheet is None:
            raise SystemExit(f"[xlsx-audit] 未找到工作表: {sheet!r}")
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def find_header(table: list[list[Any]], keywords: tuple[str, ...]) -> int:
    """返回首个命中任一关键词的行下标;找不到返回 0(按首行表头处理)。"""
    for index, row in enumerate(table[:20]):
        for cell in row:
            if isinstance(cell, str) and any(k in cell for k in keywords):
                return index
    return 0


def column_index(header_row: list[Any], names: tuple[str, ...]) -> Optional[int]:
    """在表头行里找首个含给定名字的列下标。"""
    for index, cell in enumerate(header_row):
        if isinstance(cell, str) and any(name in cell for name in names):
            return index
    return None


# ---------------------------------------------------------------------------
# 规则引擎:规则函数签名统一 (table, tolerance) -> list[dict]
# finding 形状: {rule, severity, row, detail}(row 从 1 计,0 表示表级问题)
# ---------------------------------------------------------------------------

def _finding(rule: str, severity: str, row: int, detail: str) -> dict[str, Any]:
    return {"rule": rule, "severity": severity, "row": row, "detail": detail}


def rule_total_check(table: list[list[Any]], tolerance: float) -> list[dict[str, Any]]:
    """规则一(合计校验):合计行各数值列应等于其上方数据行同列之和。"""
    findings: list[dict[str, Any]] = []
    subtotal_rows = [i for i, row in enumerate(table) if _is_subtotal_row(row)]
    if not subtotal_rows:
        return [_finding("total_check", "info", 0, "未发现 合计/总计/小计 行,跳过合计校验")]

    width = max((len(row) for row in table), default=0)
    subtotal_set = set(subtotal_rows)
    for srow in subtotal_rows:
        # 该合计行对应的覆盖范围:上一个合计行(不含)之后到本行(不含)
        start = max(
            (i for i in subtotal_rows if i < srow), default=-1
        ) + 1
        data_rows = [i for i in range(start, srow) if i not in subtotal_set]
        for col in range(1, width):
            expected_cell = to_number(table[srow][col]) if col < len(table[srow]) else None
            if expected_cell is None:
                continue
            actual = 0.0
            converted = False
            for drow in data_rows:
                value = to_number(table[drow][col]) if col < len(table[drow]) else None
                if value is not None:
                    actual += value
                    converted = True
            if not converted:
                continue  # 纯文本列不做合计校验
            if abs(actual - expected_cell) > tolerance:
                findings.append(
                    _finding(
                        "total_check",
                        "error",
                        srow + 1,
                        f"第 {srow + 1} 行第 {col + 1} 列合计不符:应为 {actual:.4g},表载 {expected_cell:.4g}",
                    )
                )
    return findings


def rule_debit_credit(table: list[list[Any]], tolerance: float) -> list[dict[str, Any]]:
    """规则二(借贷平衡):∑借方 应等于 ∑贷方(数据行,剔除合计行)。"""
    header_index = find_header(table, ("借方", "贷方", "借", "贷"))
    header_row = table[header_index] if header_index < len(table) else []
    debit_col = column_index(header_row, ("借方", "借"))
    credit_col = column_index(header_row, ("贷方", "贷"))
    if debit_col is None or credit_col is None:
        return [_finding("debit_credit", "info", 0, "未找到 借方/贷方 列,跳过借贷平衡校验")]

    debit_sum = 0.0
    credit_sum = 0.0
    saw_number = False
    for index in range(header_index + 1, len(table)):
        if _is_subtotal_row(table[index]):
            continue
        row = table[index]
        debit = to_number(row[debit_col]) if debit_col < len(row) else None
        credit = to_number(row[credit_col]) if credit_col < len(row) else None
        if debit is not None:
            debit_sum += debit
            saw_number = True
        if credit is not None:
            credit_sum += credit
            saw_number = True
    if not saw_number:
        return [_finding("debit_credit", "info", 0, "借贷列均无数值,跳过借贷平衡校验")]
    diff = debit_sum - credit_sum
    if abs(diff) > tolerance:
        return [
            _finding(
                "debit_credit",
                "error",
                0,
                f"借贷不平:∑借方={debit_sum:.4g},∑贷方={credit_sum:.4g},差额 {diff:.4g} 超容差 {tolerance}",
            )
        ]
    return []


def rule_yoy_mom(table: list[list[Any]], tolerance: float) -> list[dict[str, Any]]:
    """规则三(同比环比勾稽):同比/环比列值应与 本期/上期(或 上月)重算值一致。"""
    header_index = find_header(table, ("本期", "同比", "环比", "上月"))
    header_row = table[header_index] if header_index < len(table) else []
    current_col = column_index(header_row, ("本期",))
    prior_col = column_index(header_row, ("上期", "上年", "去年"))
    month_col = column_index(header_row, ("上月",))
    yoy_col = column_index(header_row, ("同比",))
    mom_col = column_index(header_row, ("环比",))
    if yoy_col is None and mom_col is None:
        return [_finding("yoy_mom", "info", 0, "未找到 同比/环比 列,跳过同比环比勾稽")]

    findings: list[dict[str, Any]] = []

    def check_pairs(base_col: Optional[int], pct_col: Optional[int], label: str) -> None:
        if base_col is None or pct_col is None:
            findings.append(
                _finding("yoy_mom", "info", 0, f"缺少 {label} 对比列,跳过 {label} 勾稽")
            )
            return
        for index in range(header_index + 1, len(table)):
            if _is_subtotal_row(table[index]):
                continue
            row = table[index]
            current = to_number(row[current_col]) if current_col is not None and current_col < len(row) else None
            base = to_number(row[base_col]) if base_col < len(row) else None
            pct = to_number(row[pct_col]) if pct_col < len(row) else None
            if current is None or base is None or pct is None:
                continue
            if abs(base) < 1e-12:
                findings.append(
                    _finding("yoy_mom", "error", index + 1, f"第 {index + 1} 行 {label} 分母为 0,无法勾稽但表载 {pct:.4g}")
                )
                continue
            expected = (current - base) / abs(base) * 100.0
            if abs(expected - pct) > max(tolerance * 100.0, 0.5):
                findings.append(
                    _finding(
                        "yoy_mom",
                        "error",
                        index + 1,
                        f"第 {index + 1} 行 {label} 不符:重算应为 {expected:.2f}%,表载 {pct:.2f}%",
                    )
                )

    check_pairs(prior_col, yoy_col, "同比")
    check_pairs(month_col, mom_col, "环比")
    return findings


RULES: tuple[tuple[str, Callable[[list[list[Any]], float], list[dict[str, Any]]]], ...] = (
    ("total_check", rule_total_check),
    ("debit_credit", rule_debit_credit),
    ("yoy_mom", rule_yoy_mom),
)


# ---------------------------------------------------------------------------
# CLI 入口
# ---------------------------------------------------------------------------

def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="xlsx_audit",
        description="Excel/CSV 账表勾稽审计:合计校验 / 借贷平衡 / 同比环比勾稽",
    )
    parser.add_argument("file", help=".xlsx/.xlsm/.csv 账表路径")
    parser.add_argument("--sheet", default=None, help="工作表名(仅 .xlsx)")
    parser.add_argument("--tolerance", type=float, default=TOLERANCE_DEFAULT,
                        help="金额容差,缺省 0.01")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出 findings")
    args = parser.parse_args(argv)

    try:
        table = load_table(args.file, sheet=args.sheet)
    except SystemExit as exc:
        print(exc, file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"[xlsx-audit] 文件无法读取: {exc}", file=sys.stderr)
        return 2

    if not table:
        print("[xlsx-audit] 表为空,无内容可审", file=sys.stderr)
        return 2

    findings: list[dict[str, Any]] = []
    for _name, rule_fn in RULES:
        findings.extend(rule_fn(table, args.tolerance))

    if args.json:
        print(json.dumps(
            {
                "file": args.file,
                "rows": len(table),
                "errors": sum(1 for f in findings if f["severity"] == "error"),
                "findings": findings,
            },
            ensure_ascii=False,
            indent=2,
        ))
    else:
        if not findings:
            print("[xlsx-audit] 三类勾稽规则全部通过(或无适用规则)")
        for finding in findings:
            print(f"[{finding['severity']}] {finding['rule']}: {finding['detail']}")

    return 1 if any(f["severity"] == "error" for f in findings) else 0


if __name__ == "__main__":
    sys.exit(main())
